"""Read supported local files and keep useful source information with text."""

import csv
import json
from pathlib import Path
from typing import Any

from pypdf import PdfReader


# ==========================================
# 1. CREATE A DOCUMENT RECORD
# ==========================================

def _make_record(file_name: str, text: str, **location: Any) -> dict[str, Any] | None:
    """Attach a filename and any real page, row, sheet, or JSON location."""
    cleaned_text = text.strip()
    if not cleaned_text:
        return None

    metadata = {"file_name": file_name}
    metadata.update(location)
    return {"text": cleaned_text, "metadata": metadata}


# ==========================================
# 2. LOAD PDF PAGES
# ==========================================

def _load_pdf(path: Path) -> list[dict[str, Any]]:
    """Read each PDF page separately so citations keep the real page number."""
    records = []
    reader = PdfReader(str(path))
    for page_number, page in enumerate(reader.pages, start=1):
        record = _make_record(path.name, page.extract_text() or "", page_number=page_number)
        if record:
            records.append(record)
    return records


# ==========================================
# 3. LOAD CSV ROWS
# ==========================================

def _load_csv(path: Path) -> list[dict[str, Any]]:
    """Turn each CSV row into labeled text and keep its physical row number."""
    with path.open("r", encoding="utf-8-sig", newline="") as csv_file:
        rows = list(csv.reader(csv_file))

    if not rows:
        return []

    # A single row has no obvious header, so keep its values as Column 1, etc.
    has_header = len(rows) > 1
    headers = rows[0] if has_header else []
    data_rows = rows[1:] if has_header else rows
    first_data_row_number = 2 if has_header else 1

    records = []
    for row_number, row in enumerate(data_rows, start=first_data_row_number):
        fields = []
        for column_number, value in enumerate(row, start=1):
            if value.strip():
                header = headers[column_number - 1].strip() if column_number <= len(headers) else ""
                label = header or f"Column {column_number}"
                fields.append(f"{label}: {value.strip()}")
        record = _make_record(path.name, ". ".join(fields), row_number=row_number)
        if record:
            records.append(record)
    return records


# ==========================================
# 4. LOAD EXCEL WORKSHEETS
# ==========================================

def _row_as_text(headers: list[Any], row: tuple[Any, ...]) -> str:
    """Label spreadsheet cell values with their column headings."""
    fields = []
    for column_number, value in enumerate(row, start=1):
        if value is None or str(value).strip() == "":
            continue
        header_value = headers[column_number - 1] if column_number <= len(headers) else None
        label = str(header_value).strip() if header_value is not None else ""
        fields.append(f"{label or f'Column {column_number}'}: {value}")
    return ". ".join(fields)


def _load_xlsx(path: Path) -> list[dict[str, Any]]:
    """Read .xlsx and .xlsm sheets row by row with sheet and row citations."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    records = []
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            headers = list(next(rows, ()))
            for row_number, row in enumerate(rows, start=2):
                record = _make_record(
                    path.name,
                    _row_as_text(headers, row),
                    sheet_name=sheet.title,
                    row_number=row_number,
                )
                if record:
                    records.append(record)
    finally:
        workbook.close()
    return records


def _load_xls(path: Path) -> list[dict[str, Any]]:
    """Read older .xls workbooks and keep their worksheet and row locations."""
    import xlrd

    workbook = xlrd.open_workbook(str(path), on_demand=True)
    records = []
    try:
        for sheet in workbook.sheets():
            headers = sheet.row_values(0) if sheet.nrows else []
            for row_index in range(1, sheet.nrows):
                record = _make_record(
                    path.name,
                    _row_as_text(headers, tuple(sheet.row_values(row_index))),
                    sheet_name=sheet.name,
                    row_number=row_index + 1,
                )
                if record:
                    records.append(record)
    finally:
        workbook.release_resources()
    return records


# ==========================================
# 5. LOAD JSON AND TEXT FILES
# ==========================================

def _load_json(path: Path) -> list[dict[str, Any]]:
    """Read JSON objects or arrays and preserve an identifying JSON path."""
    with path.open("r", encoding="utf-8-sig") as json_file:
        content = json.load(json_file)

    if isinstance(content, list):
        items = [(f"$[{index}]", item) for index, item in enumerate(content)]
    else:
        items = [("$", content)]

    records = []
    for json_path, item in items:
        text = json.dumps(item, ensure_ascii=False, indent=2)
        record = _make_record(path.name, text, json_path=json_path)
        if record:
            records.append(record)
    return records


def _load_text(path: Path) -> list[dict[str, Any]]:
    """Read plain text or Markdown files; the filename is their source."""
    record = _make_record(path.name, path.read_text(encoding="utf-8-sig"))
    return [record] if record else []


# ==========================================
# 6. LOAD ALL SUPPORTED DOCUMENTS
# ==========================================

def load_documents(folder_path: str | Path) -> list[dict[str, Any]]:
    """Read supported files from a folder and retain truthful source metadata.

    Supported formats are PDF, CSV, .xlsx/.xlsm, .xls, JSON, TXT, and Markdown.
    Damaged supported files are reported and skipped without stopping the
    other files. Files with unsupported extensions are ignored.
    """
    folder = Path(folder_path)
    if not folder.exists():
        return []

    loaders = {
        ".pdf": _load_pdf,
        ".csv": _load_csv,
        ".xlsx": _load_xlsx,
        ".xlsm": _load_xlsx,
        ".xls": _load_xls,
        ".json": _load_json,
        ".txt": _load_text,
        ".md": _load_text,
    }
    documents: list[dict[str, Any]] = []
    for path in sorted(folder.iterdir()):
        loader = loaders.get(path.suffix.lower())
        if not path.is_file() or loader is None:
            continue
        try:
            documents.extend(loader(path))
        except Exception as error:
            print(f"Could not read {path.name}: {error}")

    return documents
